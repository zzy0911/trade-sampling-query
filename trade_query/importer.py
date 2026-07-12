from __future__ import annotations

import math
import re
from decimal import Decimal
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook

from .database import Database, clean_text, numeric_or_none, utc_now


INDUSTRIES = {
    "51": ("批发业", "销售额"),
    "52": ("零售业", "销售额"),
    "61": ("住宿业", "营业额"),
    "62": ("餐饮业", "营业额"),
}


def industry_info(code: Any) -> tuple[str, str] | None:
    normalized = normalize_code(code)
    return INDUSTRIES.get(normalized[:2])


def normalize_code(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        if not math.isfinite(value):
            return ""
        # Excel numbers retain about 15 significant decimal digits.  Using the
        # binary float's exact integer would turn a displayed …610000 code into
        # …609984, so reconstruct the same decimal representation Excel shows.
        decimal_value = Decimal(format(value, ".15g"))
        if decimal_value == decimal_value.to_integral_value():
            return format(decimal_value.quantize(Decimal("1")), "f")
        return format(decimal_value, "f").rstrip("0").rstrip(".")
    text = str(value).strip()
    if re.fullmatch(r"\d+\.0+", text):
        return text.split(".", 1)[0]
    return text


def normalize_header(value: Any) -> str:
    return re.sub(r"\s+", "", clean_text(value)).replace("：", ":")


def extract_year(filename: str) -> int | None:
    match = re.search(r"(20\d{2})", filename)
    return int(match.group(1)) if match else None


def extract_quarter(sheet_name: str) -> int | None:
    match = re.search(r"第\s*([一二三四1-4])\s*季度", sheet_name)
    if not match:
        return None
    value = match.group(1)
    return {"一": 1, "二": 2, "三": 3, "四": 4}[value] if value in "一二三四" else int(value)


def _column_map(headers: Iterable[Any]) -> dict[str, int]:
    normalized = [normalize_header(value) for value in headers]

    def find(*terms: str) -> int | None:
        for index, header in enumerate(normalized):
            if all(term in header for term in terms):
                return index
        return None

    result = {
        "subregion": find("处理地"),
        "unit_code": find("单位代码"),
        "unit_name": find("单位名称"),
        "industry_code": find("行业代码"),
        "sales_current": find("商品销售额", "本季度"),
        "sales_previous": find("商品销售额", "上年同期"),
        "turnover_current": find("营业额", "本季度"),
        "turnover_previous": find("营业额", "上年同期"),
        "explanation": find("解释"),
    }
    required = ["subregion", "unit_code", "unit_name", "industry_code"]
    missing = [name for name in required if result[name] is None]
    if missing:
        raise ValueError(f"缺少必要列：{', '.join(missing)}")
    if result["sales_current"] is None and result["turnover_current"] is None:
        raise ValueError("未找到本季度销售额或营业额列")
    return {name: index for name, index in result.items() if index is not None}


def _cell(row: tuple[Any, ...], columns: dict[str, int], name: str) -> Any:
    index = columns.get(name)
    return None if index is None or index >= len(row) else row[index]


def read_workbook(path: str | Path, district: str, year: int) -> list[dict[str, Any]]:
    source_path = Path(path)
    workbook = load_workbook(source_path, read_only=True, data_only=True)
    records: list[dict[str, Any]] = []
    try:
        for worksheet in workbook.worksheets:
            quarter = extract_quarter(worksheet.title)
            if quarter is None:
                continue
            rows = worksheet.iter_rows(values_only=True)
            try:
                headers = next(rows)
            except StopIteration:
                continue
            columns = _column_map(headers)
            for row_number, row in enumerate(rows, start=2):
                unit_name = clean_text(_cell(row, columns, "unit_name"))
                industry_code = normalize_code(_cell(row, columns, "industry_code"))
                if not unit_name and not industry_code:
                    continue
                info = industry_info(industry_code)
                if info is None:
                    raise ValueError(
                        f"{worksheet.title} 第{row_number}行的行业代码“{industry_code}”不受支持"
                    )
                industry_name, metric_kind = info
                value_prefix = "sales" if metric_kind == "销售额" else "turnover"
                current = numeric_or_none(_cell(row, columns, f"{value_prefix}_current"))
                previous = numeric_or_none(_cell(row, columns, f"{value_prefix}_previous"))
                yoy = (
                    None
                    if previous in (None, 0) or current is None
                    else (current - previous) / previous * 100
                )
                explanation_value = _cell(row, columns, "explanation")
                explanation = clean_text(explanation_value)
                if explanation.startswith("#"):
                    explanation = ""
                records.append(
                    {
                        "district": district,
                        "year": year,
                        "quarter": quarter,
                        "subregion": clean_text(_cell(row, columns, "subregion")),
                        "unit_code": normalize_code(_cell(row, columns, "unit_code")),
                        "unit_name": unit_name,
                        "industry_code": industry_code,
                        "industry_name": industry_name,
                        "metric_kind": metric_kind,
                        "current_value": current,
                        "previous_value": previous,
                        "yoy_rate": yoy,
                        "explanation": explanation or None,
                        "source_file": source_path.name,
                    }
                )
    finally:
        workbook.close()
    if not records:
        raise ValueError("没有找到可导入的季度数据")
    return records


def import_workbook(
    database: Database, path: str | Path, district: str, year: int | None = None
) -> dict[str, Any]:
    source_path = Path(path)
    selected_year = year or extract_year(source_path.name)
    if selected_year is None:
        raise ValueError("无法从文件名识别年份，请手工指定年份")
    if not (2000 <= selected_year <= 2100):
        raise ValueError("年份必须在 2000 到 2100 之间")
    district = clean_text(district)
    if not district:
        raise ValueError("区域不能为空")

    now = utc_now()
    try:
        records = read_workbook(source_path, district, selected_year)
        with database.connect() as connection:
            quarters = sorted({record["quarter"] for record in records})
            connection.executemany(
                "DELETE FROM records WHERE district = ? AND year = ? AND quarter = ?",
                [(district, selected_year, quarter) for quarter in quarters],
            )
            connection.executemany(
                """
                INSERT INTO records(
                    district, year, quarter, subregion, unit_code, unit_name,
                    industry_code, industry_name, metric_kind, current_value,
                    previous_value, yoy_rate, explanation, source_file,
                    created_at, updated_at
                ) VALUES(
                    :district, :year, :quarter, :subregion, :unit_code, :unit_name,
                    :industry_code, :industry_name, :metric_kind, :current_value,
                    :previous_value, :yoy_rate, :explanation, :source_file,
                    :created_at, :updated_at
                )
                ON CONFLICT(district, year, quarter, unit_code, unit_name)
                DO UPDATE SET
                    subregion = excluded.subregion,
                    industry_code = excluded.industry_code,
                    industry_name = excluded.industry_name,
                    metric_kind = excluded.metric_kind,
                    current_value = excluded.current_value,
                    previous_value = excluded.previous_value,
                    yoy_rate = excluded.yoy_rate,
                    explanation = excluded.explanation,
                    source_file = excluded.source_file,
                    updated_at = excluded.updated_at
                """,
                [{**record, "created_at": now, "updated_at": now} for record in records],
            )
            cursor = connection.execute(
                """
                INSERT INTO import_batches(
                    district, year, source_file, imported_rows, status, message, created_at
                ) VALUES(?, ?, ?, ?, 'success', ?, ?)
                """,
                (
                    district,
                    selected_year,
                    source_path.name,
                    len(records),
                    "导入或更新完成",
                    now,
                ),
            )
            batch_id = cursor.lastrowid
        return {
            "batch_id": batch_id,
            "district": district,
            "year": selected_year,
            "source_file": source_path.name,
            "imported_rows": len(records),
        }
    except Exception as exc:
        with database.connect() as connection:
            connection.execute(
                """
                INSERT INTO import_batches(
                    district, year, source_file, imported_rows, status, message, created_at
                ) VALUES(?, ?, ?, 0, 'failed', ?, ?)
                """,
                (district, selected_year, source_path.name, str(exc), now),
            )
        raise
