from __future__ import annotations

import tempfile
import threading
import unittest
import urllib.request
from email.parser import BytesParser
from email.policy import default
from pathlib import Path
from http.server import ThreadingHTTPServer

from openpyxl import Workbook

from trade_query.database import Database, hash_password, normalize_subregion, verify_password
from trade_query.importer import (
    DuplicateSourceFileError,
    extract_quarter,
    import_workbook,
    industry_info,
    normalize_code,
    read_workbook,
)
from trade_query.server import decode_multipart_text
from trade_query import server as server_module


class CoreTests(unittest.TestCase):
    def test_frontend_avoids_array_at_for_legacy_browsers(self) -> None:
        static_dir = Path(__file__).resolve().parent.parent / "static"
        for script in static_dir.glob("*.js"):
            self.assertNotIn(".at(", script.read_text(encoding="utf-8"), script.name)

    def test_frontend_contains_detail_filters_editing_and_import_dialog(self) -> None:
        static_dir = Path(__file__).resolve().parent.parent / "static"
        history_html = (static_dir / "history.html").read_text(encoding="utf-8")
        history_js = (static_dir / "history.js").read_text(encoding="utf-8")
        admin_html = (static_dir / "admin.html").read_text(encoding="utf-8")
        self.assertIn('id="subregion"', history_html)
        self.assertIn('id="replacement"', history_html)
        self.assertNotIn("蓝色为新替换单位", history_html)
        self.assertIn("public-explanation", history_js)
        self.assertIn("/explanation", history_js)
        self.assertIn("已导入同名数据，确认后将覆盖原数据", admin_html)

    def test_industry_mapping(self) -> None:
        self.assertEqual(industry_info(5123), ("批发业", "销售额"))
        self.assertEqual(industry_info("5299"), ("零售业", "销售额"))
        self.assertEqual(industry_info(6129), ("住宿业", "营业额"))
        self.assertEqual(industry_info("6291"), ("餐饮业", "营业额"))
        self.assertIsNone(industry_info("7010"))

    def test_normalize_long_code(self) -> None:
        self.assertEqual(normalize_code(320111003010510000), "320111003010510000")
        self.assertEqual(normalize_code(3.2011100400261e17), "320111004002610000")
        self.assertEqual(normalize_code("91320191MA1Y53AC7W"), "91320191MA1Y53AC7W")

    def test_normalize_subregion_removes_only_trailing_area_code(self) -> None:
        self.assertEqual(normalize_subregion("江浦街道(320111004000)"), "江浦街道")
        self.assertEqual(normalize_subregion("江浦街道（320111004001）"), "江浦街道")
        self.assertEqual(normalize_subregion("江浦街道"), "江浦街道")
        self.assertEqual(normalize_subregion("测试街道（北区）"), "测试街道（北区）")

    def test_quarter_parsing(self) -> None:
        self.assertEqual(extract_quarter("第1季度"), 1)
        self.assertEqual(extract_quarter("第二季度"), 2)
        self.assertIsNone(extract_quarter("汇总"))

    def test_previous_value_comes_from_the_same_quarter_sheet(self) -> None:
        headers = [
            "处理地名称",
            "单位代码",
            "单位名称",
            "行业代码",
            "商品销售额（批发和零售业单位填报）;本季度",
            "商品销售额（批发和零售业单位填报）;上年同期",
        ]
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "测试2026.xlsx"
            workbook = Workbook()
            first_quarter = workbook.active
            first_quarter.title = "第1季度"
            first_quarter.append(headers)
            first_quarter.append(["街道", "001", "测试单位", "5219", 100, 10])
            second_quarter = workbook.create_sheet("第2季度")
            second_quarter.append(headers)
            second_quarter.append(["街道", "001", "测试单位", "5219", 220, 40])
            workbook.save(path)
            workbook.close()

            records = read_workbook(path, "测试区域", 2026)
            second_record = next(record for record in records if record["quarter"] == 2)
            self.assertEqual(second_record["current_value"], 220)
            self.assertEqual(second_record["previous_value"], 40)
            self.assertAlmostEqual(second_record["yoy_rate"], 450)

    def test_import_normalizes_subregion_name(self) -> None:
        headers = [
            "处理地名称",
            "单位代码",
            "单位名称",
            "行业代码",
            "商品销售额（批发和零售业单位填报）;本季度",
        ]
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "测试2026.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "第1季度"
            sheet.append(headers)
            sheet.append(["江浦街道(320111004001)", "001", "测试单位", "5219", 100])
            workbook.save(path)
            workbook.close()

            records = read_workbook(path, "测试区域", 2026)
            self.assertEqual(records[0]["subregion"], "江浦街道")

    def test_import_confirmation_merge_and_full_source_replacement(self) -> None:
        headers = [
            "处理地名称",
            "单位代码",
            "单位名称",
            "行业代码",
            "商品销售额（批发和零售业单位填报）;本季度",
            "商品销售额（批发和零售业单位填报）;上年同期",
        ]
        with tempfile.TemporaryDirectory() as directory:
            directory_path = Path(directory)
            database = Database(directory_path / "test.db")
            database.initialize()
            source_a = directory_path / "A2026.xlsx"
            workbook = Workbook()
            first = workbook.active
            first.title = "第1季度"
            first.append(headers)
            first.append(["街道甲", "001", "A一季度", "5219", 100, 90])
            second = workbook.create_sheet("第2季度")
            second.append(headers)
            second.append(["街道甲", "002", "A二季度", "5219", 200, 180])
            workbook.save(source_a)
            workbook.close()

            first_result = import_workbook(database, source_a, "测试区域")
            self.assertFalse(first_result["overwritten"])
            with self.assertRaises(DuplicateSourceFileError):
                import_workbook(database, source_a, "测试区域")

            source_b = directory_path / "B2026.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "第1季度"
            sheet.append(headers)
            sheet.append(["街道乙", "003", "B一季度", "5219", 300, 270])
            workbook.save(source_b)
            workbook.close()
            import_workbook(database, source_b, "测试区域")

            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "第2季度"
            sheet.append(headers)
            sheet.append(["街道甲", "004", "A新版", "5219", 400, 360])
            workbook.save(source_a)
            workbook.close()
            replaced = import_workbook(database, source_a, "测试区域", overwrite=True)
            self.assertTrue(replaced["overwritten"])
            with database.connect() as connection:
                rows = connection.execute(
                    "SELECT source_file, quarter, unit_name FROM records ORDER BY source_file"
                ).fetchall()
            self.assertEqual(
                [(row["source_file"], row["quarter"], row["unit_name"]) for row in rows],
                [("A2026.xlsx", 2, "A新版"), ("B2026.xlsx", 1, "B一季度")],
            )

    def test_invalid_replacement_does_not_delete_existing_source(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            directory_path = Path(directory)
            database = Database(directory_path / "test.db")
            database.initialize()
            source = directory_path / "source2026.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "第1季度"
            sheet.append(["处理地名称", "单位代码", "单位名称", "行业代码", "商品销售额;本季度"])
            sheet.append(["街道", "001", "原数据", "5219", 100])
            workbook.save(source)
            workbook.close()
            import_workbook(database, source, "测试区域")

            workbook = Workbook()
            workbook.active.title = "无效工作表"
            workbook.save(source)
            workbook.close()
            with self.assertRaises(ValueError):
                import_workbook(database, source, "测试区域", overwrite=True)
            with database.connect() as connection:
                row = connection.execute("SELECT unit_name FROM records").fetchone()
            self.assertEqual(row["unit_name"], "原数据")

    def test_password_hash(self) -> None:
        encoded = hash_password("correct-password")
        self.assertTrue(verify_password("correct-password", encoded))
        self.assertFalse(verify_password("wrong-password", encoded))

    def test_browser_formdata_chinese_field(self) -> None:
        boundary = "test-boundary"
        body = (
            f"Content-Type: multipart/form-data; boundary={boundary}\r\n"
            "MIME-Version: 1.0\r\n\r\n"
            f"--{boundary}\r\n"
            'Content-Disposition: form-data; name="district"\r\n\r\n'
        ).encode() + "江北新区".encode("utf-8") + f"\r\n--{boundary}--\r\n".encode()
        message = BytesParser(policy=default).parsebytes(body)
        part = next(message.iter_parts())
        self.assertEqual(decode_multipart_text(part), "江北新区")

    def test_database_aggregate_and_update(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            database = Database(Path(directory) / "test.db")
            database.initialize()
            now = "2026-01-01T00:00:00+00:00"
            with database.connect() as connection:
                connection.execute(
                    """
                    INSERT INTO records(
                        district, year, quarter, subregion, unit_code, unit_name,
                        industry_code, industry_name, metric_kind, current_value,
                        previous_value, yoy_rate, explanation, source_file, created_at, updated_at
                    ) VALUES('江北新区', 2026, 1, '沿江街道', '001', '测试单位',
                             '5219', '零售业', '销售额', 120, 100, 20, NULL, 'test.xlsx', ?, ?)
                    """,
                    (now, now),
                )
            trend = database.trend("all", "all")
            self.assertEqual(len(trend), 1)
            self.assertAlmostEqual(trend[0]["yoy_rate"], 20)
            record = database.records("all", "all", 2026, 1)[0]
            updated = database.update_record(
                record["id"], {"current_value": 90, "previous_value": 100, "explanation": "修订"}
            )
            self.assertAlmostEqual(updated["yoy_rate"], -10)
            self.assertEqual(updated["explanation"], "修订")
            with database.connect() as connection:
                connection.execute(
                    "UPDATE records SET source_file = ? WHERE id = ?",
                    ("JBXQ2026\ufffd\ufffd\ufffd\ufffd.xlsx", record["id"]),
                )
            self.assertEqual(
                database.matching_source_files("JBXQ2026年数据.xlsx"),
                ["JBXQ2026\ufffd\ufffd\ufffd\ufffd.xlsx"],
            )
            self.assertEqual(database.matching_source_files("PKQ2026年数据.xlsx"), [])

    def test_legacy_subregion_codes_are_deduplicated_and_filterable(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            database = Database(Path(directory) / "test.db")
            database.initialize()
            now = "2026-01-01T00:00:00+00:00"
            rows = [
                ("江浦街道(320111004001)", "001", "测试单位甲"),
                ("江浦街道（320111004002）", "002", "测试单位乙"),
            ]
            with database.connect() as connection:
                connection.executemany(
                    """
                    INSERT INTO records(
                        district, year, quarter, subregion, unit_code, unit_name,
                        industry_code, industry_name, metric_kind, current_value,
                        previous_value, yoy_rate, explanation, source_file, created_at, updated_at
                    ) VALUES('浦口区', 2026, 1, ?, ?, ?,
                             '5219', '零售业', '销售额', 120, 100, 20, NULL, 'test.xlsx', ?, ?)
                    """,
                    [(*row, now, now) for row in rows],
                )

            self.assertEqual(
                database.subregions("浦口区", "all", 2026, 1), ["江浦街道"]
            )
            records = database.records(
                "浦口区", "all", 2026, 1, subregion="江浦街道"
            )
            self.assertEqual(len(records), 2)
            self.assertEqual({record["subregion"] for record in records}, {"江浦街道"})
            self.assertEqual(
                database.record_count(
                    "浦口区", "all", 2026, 1, subregion="江浦街道"
                ),
                2,
            )

    def test_public_explanation_http_endpoint_needs_no_login(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            database = Database(Path(directory) / "test.db")
            database.initialize()
            now = "2026-01-01T00:00:00+00:00"
            with database.connect() as connection:
                cursor = connection.execute(
                    """
                    INSERT INTO records(
                        district, year, quarter, subregion, unit_code, unit_name,
                        industry_code, industry_name, metric_kind, current_value,
                        previous_value, yoy_rate, explanation, source_file, created_at, updated_at
                    ) VALUES('测试区域', 2026, 1, '测试街道', '001', '测试单位',
                             '5219', '零售业', '销售额', 120, 100, 20, NULL, 'test.xlsx', ?, ?)
                    """,
                    (now, now),
                )
                record_id = cursor.lastrowid

            class QuietHandler(server_module.AppHandler):
                def log_message(self, format: str, *args: object) -> None:
                    return

            original_database = server_module.DB
            server_module.DB = database
            httpd = ThreadingHTTPServer(("127.0.0.1", 0), QuietHandler)
            thread = threading.Thread(target=httpd.serve_forever, daemon=True)
            thread.start()
            try:
                request = urllib.request.Request(
                    f"http://127.0.0.1:{httpd.server_port}/api/records/{record_id}/explanation",
                    data=b'{"explanation":"visitor edit"}',
                    headers={"Content-Type": "application/json"},
                    method="PUT",
                )
                with urllib.request.urlopen(request) as response:
                    payload = response.read().decode("utf-8")
                self.assertIn("visitor edit", payload)
                with database.connect() as connection:
                    saved = connection.execute(
                        "SELECT explanation FROM records WHERE id = ?", (record_id,)
                    ).fetchone()
                self.assertEqual(saved["explanation"], "visitor edit")
            finally:
                httpd.shutdown()
                httpd.server_close()
                thread.join(timeout=2)
                server_module.DB = original_database

    def test_records_mark_units_replaced_since_previous_quarter(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            database = Database(Path(directory) / "test.db")
            database.initialize()
            now = "2026-01-01T00:00:00+00:00"
            rows = [
                (2025, 4, "测试街道甲", "001", "延续单位"),
                (2026, 1, "测试街道甲", "001", "延续单位"),
                (2026, 1, "测试街道乙", "002", "新替换单位"),
            ]
            with database.connect() as connection:
                connection.executemany(
                    """
                    INSERT INTO records(
                        district, year, quarter, subregion, unit_code, unit_name,
                        industry_code, industry_name, metric_kind, current_value,
                        previous_value, yoy_rate, explanation, source_file, created_at, updated_at
                    ) VALUES('测试区域', ?, ?, ?, ?, ?,
                             '5219', '零售业', '销售额', 120, 100, 20, NULL, 'test.xlsx', ?, ?)
                    """,
                    [
                        (year, quarter, subregion, code, name, now, now)
                        for year, quarter, subregion, code, name in rows
                    ],
                )

            records = database.records("测试区域", "all", 2026, 1)
            replacement_flags = {record["unit_name"]: record["is_new_unit"] for record in records}
            self.assertEqual(replacement_flags["延续单位"], 0)
            self.assertEqual(replacement_flags["新替换单位"], 1)
            self.assertEqual(
                set(database.subregions("测试区域", "all", 2026, 1)),
                {"测试街道甲", "测试街道乙"},
            )
            new_records = database.records(
                "测试区域", "all", 2026, 1, subregion="测试街道乙", replacement="new"
            )
            self.assertEqual([record["unit_name"] for record in new_records], ["新替换单位"])
            self.assertEqual(
                database.record_count(
                    "测试区域", "all", 2026, 1, replacement="existing"
                ),
                1,
            )
            updated = database.update_explanation(new_records[0]["id"], "  访客修订  ")
            self.assertEqual(updated["explanation"], "访客修订")


if __name__ == "__main__":
    unittest.main()
